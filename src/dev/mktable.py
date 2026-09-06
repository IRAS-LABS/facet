"""
Fixtures for the table viewer, plus the truth to check it against.

Every file here is written by something that is not FACET — duckdb for
Parquet, openpyxl for Excel, the stdlib csv module for delimited text — and
each one also dumps what *it* thinks the contents are. That is the whole
point: a reader tested against fixtures it generated itself proves only that
it is self-consistent.
"""
import csv, json, os, random, sys
import duckdb
from openpyxl import Workbook
from datetime import datetime, date

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tbl")
os.makedirs(OUT, exist_ok=True)
truth = {}

# ── CSV: the awkward cases, not the easy one ────────────────────────────────
rows = [
    ["id", "name", "note", "amount", "when"],
    ["1", "Plain", "nothing special", "12.5", "2024-01-02"],
    ["2", 'Comma, inside', 'He said "hi"', "1,234.00", "2024-02-03"],
    ["3", "Newline\ninside", "line one\nline two", "-7", "2024-03-04"],
    ["4", "", "empty name", "0", ""],
    ["5", "Quote\"mid", "unquoted quote stays", "99.999", "2024-05-06"],
    ["6", "Ünïcodé ✓", "utf-8 beyond ascii", "3.14159", "2024-06-07"],
]
with open(f"{OUT}/awkward.csv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(rows)
truth["awkward.csv"] = rows

# CRLF + BOM + semicolons, i.e. what Excel writes in half of Europe.
# newline="" is not optional: csv.writer emits \r\n itself, and letting Python
# translate the \n as well produces \r\r\n, which is a different file.
with open(f"{OUT}/euro.csv", "w", newline="", encoding="utf-8-sig") as f:
    csv.writer(f, delimiter=";").writerows(rows)
truth["euro.csv"] = rows

# Tabs, with a prose column full of commas — the delimiter sniffer's real test
tsv = [["id", "sentence", "n"]] + [
    [str(i), f"one, two, three, and {i} more", str(i * 7)] for i in range(1, 30)
]
with open(f"{OUT}/prose.tsv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f, delimiter="\t").writerows(tsv)
truth["prose.tsv"] = tsv

# No header: every row is numbers
nohdr = [[str(i), str(i * 2), str(i * i)] for i in range(1, 40)]
with open(f"{OUT}/nohdr.csv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(nohdr)
truth["nohdr.csv"] = nohdr

# UTF-16LE, which is what "Unicode Text (*.txt)" means in Excel
with open(f"{OUT}/utf16.csv", "w", newline="", encoding="utf-16") as f:
    csv.writer(f).writerows(rows)
truth["utf16.csv"] = rows

# Big: 200k rows, so "must not be loaded whole" is a claim with teeth
random.seed(7)
big_path = f"{OUT}/big.csv"
with open(big_path, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["row", "uuidish", "value", "quoted"])
    for i in range(200000):
        w.writerow([i, f"{random.getrandbits(64):016x}", round(random.random() * 1000, 4),
                    'has, comma' if i % 1000 == 0 else "plain"])
with open(big_path, newline="", encoding="utf-8") as f:
    r = list(csv.reader(f))
truth["big.csv"] = {"rows": len(r), "first": r[0], "row1": r[1], "last": r[-1],
                    "row100000": r[100000], "rowsWithComma": r[1001],
                    "size": os.path.getsize(big_path)}

# ── XLSX ────────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Data"
ws.append(["id", "text", "number", "when", "flag", "formula"])
for i in range(1, 51):
    ws.append([i, f"row {i}", i * 1.5, date(2024, 1, 1 + (i % 28)), i % 2 == 0, None])
ws["F2"] = "=A2*2"
ws2 = wb.create_sheet("Second tab")
ws2.append(["only", "two", "columns"])
ws2.append(["a", "b", "c"])
ws3 = wb.create_sheet("Sparse")
ws3["A1"] = "left"
ws3["D1"] = "far right"
ws3["B5"] = "row five"          # rows 2-4 are absent from the XML entirely
ws3["C7"] = 'quote " and <tag>'  # forces XML escaping
wb.save(f"{OUT}/book.xlsx")

from openpyxl import load_workbook
lw = load_workbook(f"{OUT}/book.xlsx")
truth["book.xlsx"] = {
    "sheets": lw.sheetnames,
    "data": [[("" if c is None else
               (c.date() if isinstance(c, datetime) and c.time().isoformat() == "00:00:00" else c).isoformat()
               if isinstance(c, (date, datetime)) else
               "TRUE" if c is True else "FALSE" if c is False else str(c))
              for c in row]
             for row in lw["Data"].iter_rows(values_only=True)],
    "sparse": [[("" if c is None else str(c)) for c in row]
               for row in lw["Sparse"].iter_rows(values_only=True)],
}

# ── Parquet, written by duckdb, in several shapes ───────────────────────────
con = duckdb.connect()
con.execute("""
CREATE TABLE t AS SELECT
  i::INTEGER              AS id,
  ('name ' || i)          AS name,
  CASE WHEN i % 7 = 0 THEN NULL ELSE (i * 1.5)::DOUBLE END AS score,
  (i % 3 = 0)             AS flag,
  (DATE '2020-01-01' + i::INTEGER) AS day,
  (TIMESTAMP '2020-01-01 00:00:00' + INTERVAL (i::INTEGER) MINUTE) AS ts,
  (i * 1000000000)::BIGINT AS big,
  CASE WHEN i % 5 = 0 THEN NULL ELSE ('cat' || (i % 4)) END AS category,
  (i / 100.0)::DECIMAL(18,4) AS money
FROM range(1, 25001) t(i);
""")

variants = {
    "snappy.parquet": "(FORMAT PARQUET, COMPRESSION SNAPPY, ROW_GROUP_SIZE 5000)",
    "gzip.parquet": "(FORMAT PARQUET, COMPRESSION GZIP, ROW_GROUP_SIZE 12000)",
    "plain.parquet": "(FORMAT PARQUET, COMPRESSION UNCOMPRESSED, ROW_GROUP_SIZE 25000)",
    "v2.parquet": "(FORMAT PARQUET, COMPRESSION SNAPPY, PARQUET_VERSION V2, ROW_GROUP_SIZE 8000)",
}
for name, opts in variants.items():
    con.execute(f"COPY t TO '{OUT}/{name}' {opts}")

cols = [d[0] for d in con.execute("SELECT * FROM t LIMIT 0").description]


def cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def sample(rowset):
    return [[cell(v) for v in r] for r in rowset]


spots = [0, 1, 4999, 5000, 12000, 24999]
truth["parquet"] = {
    "columns": cols,
    "rows": con.execute("SELECT count(*) FROM t").fetchone()[0],
    "at": {str(s): sample(con.execute(
        f"SELECT * FROM t ORDER BY id LIMIT 1 OFFSET {s}").fetchall())[0] for s in spots},
    "head": sample(con.execute("SELECT * FROM t ORDER BY id LIMIT 20").fetchall()),
    "nulls_in_score": con.execute("SELECT count(*) FROM t WHERE score IS NULL").fetchone()[0],
    "nulls_in_category": con.execute("SELECT count(*) FROM t WHERE category IS NULL").fetchone()[0],
}

json.dump(truth, open(f"{OUT}/truth.json", "w", encoding="utf8"), indent=1, default=str)
print("wrote", len(os.listdir(OUT)), "files to", OUT)
for f in sorted(os.listdir(OUT)):
    print(f"  {f:16s} {os.path.getsize(os.path.join(OUT, f)):>10,}")
